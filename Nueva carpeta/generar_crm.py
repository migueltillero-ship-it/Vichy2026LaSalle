import imaplib
import email
import pandas as pd
import sys
from openpyxl.styles import Font, PatternFill, Alignment

if sys.stdout.encoding != 'utf-8':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

EMAIL_USER = "migueltillero@gmail.com"
EMAIL_PASS = "fluzxfrdilgxuhjr" 
IMAP_SERVER = "imap.gmail.com"
CRITERIO_BUSQUEDA = 'SUBJECT "Vichy 2026"'

def ejecutar_crm():
    print(f"🔄 Actualizando CRM con formato institucional...")
    try:
        mail = imaplib.IMAP4_SSL(IMAP_SERVER)
        mail.login(EMAIL_USER, EMAIL_PASS)
        mail.select("inbox")
        status, mensajes = mail.search(None, CRITERIO_BUSQUEDA)
        ids = mensajes[0].split()
        
        lista = []
        for num in ids:
            _, data = mail.fetch(num, "(RFC822)")
            msg = email.message_from_bytes(data[0][1])
            body = ""
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        body = part.get_payload(decode=True).decode('utf-8', errors='ignore')
            else:
                body = msg.get_payload(decode=True).decode('utf-8', errors='ignore')

            d = {}
            for l in body.split('\n'):
                if ":" not in l: continue
                clave, valor = [x.strip() for x in l.split(':', 1)]
                clave_low = clave.lower()
                if "nom" in clave_low: d['Nombre'] = valor
                elif "mail" in clave_low: d['Email'] = valor
                elif "tel" in clave_low or "phone" in clave_low: d['WhatsApp'] = valor
            if d: lista.append(d)
        
        if lista:
            nombre_archivo = "Base_Datos_Vichy2026.xlsx"
            df = pd.DataFrame(lista)
            writer = pd.ExcelWriter(nombre_archivo, engine='openpyxl')
            df.to_excel(writer, index=False, sheet_name='Interesados')
            
            # --- ESTILO ALLIANCE FRANÇAISE ---
            workbook = writer.book
            worksheet = writer.sheets['Interesados']
            azul_af = PatternFill(start_color="002395", end_color="002395", fill_type="solid")
            blanco = Font(color="FFFFFF", bold=True)
            
            for col_num, value in enumerate(df.columns.values):
                cell = worksheet.cell(row=1, column=col_num + 1)
                cell.fill = azul_af
                cell.font = blanco
                cell.alignment = Alignment(horizontal="center")
                worksheet.column_dimensions[chr(65 + col_num)].width = 30

            writer.close()
            print(f"✅ CRM ACTUALIZADO: '{nombre_archivo}' ahora tiene formato de lujo.")
        else:
            print("❌ No hay datos nuevos.")
    except Exception as e:
        print(f"❌ ERROR: {e}")

if __name__ == "__main__":
    ejecutar_crm()